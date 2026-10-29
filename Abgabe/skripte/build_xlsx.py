# -*- coding: utf-8 -*-
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
import data as D

OUT = sys.argv[1]
RAW = "/home/user/ERP-Systeme/KVA_Fallpaket_Studierende_1/07_Daten/Anforderungskatalog_ERP2028_Rohfassung.xlsx"

NAVY = "1F4E79"
HFONT = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor=NAVY)
SUBFILL = PatternFill("solid", fgColor="DDEBF7")
INFILL = PatternFill("solid", fgColor="FFF2CC")  # Eingabe/Annahme
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
PRIO_FILL = {"K.o.": "F8CBAD", "Muss": "FCE4D6", "Kann": "E2EFDA", "Phase 2": "EDEDED", "gestrichen": "D9D9D9"}

wb = openpyxl.Workbook()


def sheet(title, headers, widths, rows, freeze="A2"):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment, c.border = HFONT, HFILL, Alignment(wrap_text=True, vertical="center"), BORDER
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border = WRAP, BORDER
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    return ws


# ------------------------------------------------------------ 0 Lesehilfe
ws = wb.active
ws.title = "0_Lesehilfe"
lines = [
    ("KVA – ERP 2028: Konsolidierter Anforderungskatalog, Widerspruchsliste, Bewertung & TCO", True),
    ("Artefakte B + C · HTWG Konstanz · ERP-Systeme · Stand: 29.10.2026", False),
    ("", False),
    ("Blätter", True),
    ("1_Widersprüche – 42 Widersprüche/Lücken mit Fundstelle, Auswirkung, Lösungsvorschlag und Entscheider", False),
    ("2_K.o.-Kriterien – 8 harte Ausschlusskriterien mit Prüfmethode", False),
    ("3_Anforderungen – alle 111 Rohanforderungen klassifiziert + 8 neue Anforderungen (N-xxx)", False),
    ("4_Mengengerüst – widersprüchliche Zahlen und gewählter Planwert", False),
    ("5_Longlist – 16 Optionen, Filter über K.o.-Kriterien → Shortlist", False),
    ("6_Nutzwertanalyse – gewichtete Bewertung der Shortlist + Sensitivitätsanalyse (Formeln)", False),
    ("7_TCO_Annahmen – alle Kostenparameter (gelb = Annahme, änderbar)", False),
    ("8_TCO_5J – 5-Jahres-TCO 2027–2031 je System (Formeln auf Blatt 7)", False),
    ("9_Nutzen – Business-Case-Hebel (Spannen)", False),
    ("10_Quellen – Fallpaket-Fundstellen und externe Quellen", False),
    ("11_Störinformationen – 20 bewusst aussortierte Informationen (veraltet, falsch, Anekdote, Präferenz)", False),
    ("12_Lizenzmodelle – Lizenzmodelle der Kandidaten und Folgen für KVA", False),
    ("13_Länder_Recht – DE (Inland), CZ (EU-Ausland), CH (Drittland), Datenstandort", False),
    ("15_Testskript – 6 identische Testfälle für Drehbuch-Demo und PoC aller Anbieter", False),
    ("14_Stakeholder – Anliegen von GF, CFO, IT, Pilsen, Betriebsrat, Beirat und unsere Antwort", False),
    ("", False),
    ("Legende Priorität", True),
    ("K.o. = Ausschluss bei Nichterfüllung · Muss = Phase 1 erforderlich · Kann = Bonus in Bewertung · Phase 2 = nach Go-live · gestrichen = Lösungsvorgabe/Dublette/Präferenz", False),
    ("Klassifizierung: Anforderung · gesetzlich · Rahmenbedingung · Lösungsvorgabe (beschreibt eine Lösung statt Bedarf) · Präferenz (persönlich) · Dublette · veraltet", False),
    ("", False),
    ("Wichtig: Preise von proALPHA, IFS und abas sind nicht öffentlich – alle TCO-Werte sind gekennzeichnete Annahmen (gelb) und im RFP zu validieren.", True),
]
for t, b in lines:
    ws.append([t])
    ws.cell(ws.max_row, 1).font = Font(bold=b, size=14 if ws.max_row == 1 else 11, color=NAVY if b else "000000")
ws.column_dimensions["A"].width = 140

# ------------------------------------------------------------ 1 Widersprüche
sheet("1_Widersprüche",
      ["ID", "Thema", "Aussage A (Fundstelle)", "Aussage B (Fundstelle)", "Typ", "Auswirkung", "Lösungsvorschlag", "Entscheidung durch", "Relevanz"],
      [7, 18, 45, 45, 16, 32, 55, 13, 10], D.WIDERSPRUECHE, freeze="C2")

# ------------------------------------------------------------ 2 K.o.
sheet("2_K.o.-Kriterien", ["ID", "K.o.-Kriterium", "Definition", "Bezug Anforderungen", "Prüfmethode"],
      [6, 38, 60, 24, 40], D.KO)

# ------------------------------------------------------------ 3 Anforderungen
raw = openpyxl.load_workbook(RAW)["Anforderungen"]
rows = []
for r in raw.iter_rows(min_row=2, values_only=True):
    rid, ber, tit, bes, prio, quelle, ap, status, kom = r
    k, neu, begr, ref = D.KLASS[rid]
    rows.append((rid, ber, tit, bes, prio, quelle, status, k, neu, begr, ref))
for n in D.NEU:
    rid, ber, tit, bes, k, neu, begr, ref = n
    rows.append((rid, ber, tit, bes, "–", "HTWG-Analyse", "neu", k, neu, begr, ref))
ws = sheet("3_Anforderungen",
           ["ID", "Bereich", "Titel", "Beschreibung (Rohfassung)", "Prio Fachbereich (alt)", "Quelle", "Status alt",
            "Klassifizierung", "Prio konsolidiert", "Begründung / Präzisierung", "Bezug Widerspruch"],
           [7, 11, 26, 42, 12, 16, 11, 20, 18, 50, 12], rows, freeze="D2")
for row in ws.iter_rows(min_row=2):
    p = str(row[8].value)
    for key, col in PRIO_FILL.items():
        if p.startswith(key):
            row[8].fill = PatternFill("solid", fgColor=col)
            break
# Auswertung
n0 = ws.max_row + 2
ws.cell(n0, 1, "Auswertung").font = Font(bold=True, color=NAVY)
for j, key in enumerate(["K.o.", "Muss*", "Kann", "Phase 2*", "gestrichen", "Bewertungskriterium", "GF-Entscheidung"]):
    ws.cell(n0 + 1 + j, 1, key.replace("*", " (inkl. umformuliert/Teil)"))
    ws.cell(n0 + 1 + j, 2, f'=COUNTIF($I$2:$I${len(rows)+1},"{key}")')

# ------------------------------------------------------------ 4 Mengengerüst
sheet("4_Mengengerüst", ["Größe", "Werte im Fallpaket (Quellen)", "Planwert", "Begründung", "Bezug"],
      [26, 60, 18, 55, 9], D.MENGEN)

# ------------------------------------------------------------ 5 Longlist
ws = sheet("5_Longlist",
           ["Nr", "Anbieter", "Produkt", "Herkunft", "Betriebsmodelle", "Fit CTO/ETO", "Präsenz CZ", "Partner DE",
            "Ergebnis", "Begründung", "Quelle", "Prüfstatus (Abruf 10/2026)"],
           [4, 16, 30, 11, 26, 30, 30, 11, 16, 50, 34, 24], D.LONGLIST, freeze="D2")
for row in ws.iter_rows(min_row=2):
    v = str(row[8].value)
    row[8].fill = PatternFill("solid", fgColor="C6EFCE" if v.startswith("Shortlist") else "FFEB9C" if v.startswith("Nachr") else "FFC7CE")
n0 = ws.max_row + 2
ws.cell(n0, 1, "Filterlogik").font = Font(bold=True, color=NAVY)
for j, t in enumerate([
        "Filter 1 – Zielgruppe: ERP für diskrete Fertiger im Mittelstand (200–1.000 MA), DACH-Marktpräsenz",
        "Filter 2 – K.o. K1/K2: Variantenkonfiguration + Projektfertigung im Standard oder über zertifizierte Erweiterung",
        "Filter 3 – K.o. K4/K6: Lokalisierung CZ + Implementierungspartner in DE UND CZ belegt",
        "Filter 4 – Wirtschaftlichkeit/Akzeptanz: TCO-Größenordnung passend, Betriebsmodell mit Plan B, Lehren aus 2016"]):
    ws.cell(n0 + 1 + j, 1, t)

# ------------------------------------------------------------ 6 NWA
ws = wb.create_sheet("6_Nutzwertanalyse")
NS = len(D.SYSTEME)
PC = [get_column_letter(4 + j) for j in range(NS)]            # Punkte-Spalten
GC = [get_column_letter(4 + NS + j) for j in range(NS)]       # gewichtet
TC = [get_column_letter(3 + k) for k in D.NWA_TO_TCO]         # Spalten in Blatt 8
hdr = ["Nr", "Kriterium", "Gewicht %"] + [f"Punkte {s}" for s in D.SYSTEME] + [f"Gewichtet {s}" for s in D.SYSTEME] + ["Begründung"]
ws.append(hdr)
for i, (k, g, sc, b) in enumerate(D.NWA, 1):
    r = i + 1
    ws.append([i, k, g, *sc, *[f"=$C{r}*{pc}{r}/100" for pc in PC], b])
last = len(D.NWA) + 1
blank = [""] * NS
ws.append(["", "Summe / Nutzwert (max. 5,0)", f"=SUM(C2:C{last})", *blank, *[f"=SUM({gc}2:{gc}{last})" for gc in GC], ""])
tot = last + 1
rng = f"${GC[0]}${tot}:${GC[-1]}${tot}"
ws.append(["", "Rang", "", *blank, *[f"=RANK({gc}{tot},{rng})" for gc in GC], ""])
ws.append([])
ws.append(["", "TCO 5 Jahre (Mio. €, aus Blatt 8)", "", *blank, *[f"='8_TCO_5J'!{tc}9/1000000" for tc in TC], "BC-Basis-Variante nicht in NWA (reduzierter Scope)"])
ws.append(["", "Nutzwert je Mio. € TCO", "", *blank, *[f"={gc}{tot}/{gc}{tot+3}" for gc in GC], "Kosten bewusst NICHT im Nutzwert – getrennte Betrachtung"])
for c in ws[1]:
    c.font, c.fill, c.alignment, c.border = HFONT, HFILL, Alignment(wrap_text=True, vertical="center"), BORDER
for row in ws.iter_rows(min_row=2, max_row=tot + 4):
    for c in row:
        c.alignment, c.border = WRAP, BORDER
for r in range(2, last + 1):
    for col in ["C"] + PC:
        ws[f"{col}{r}"].fill = INFILL
for r in (tot, tot + 1, tot + 3, tot + 4):
    for c in ws[r]:
        c.font = Font(bold=True)
        c.fill = SUBFILL
for r in range(2, tot + 5):
    for col in GC:
        ws[f"{col}{r}"].number_format = "0.00"
for i, w in enumerate([4, 52, 10] + [11] * NS + [12] * NS + [60], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "C2"
s0 = tot + 6
ws.cell(s0, 2, "Skala: 1 = nicht/kaum erfüllt · 2 = mit Individualentwicklung · 3 = mit Partner-Erweiterung/Einschränkung · 4 = Standard, gut · 5 = Standard, Marktstärke").font = Font(italic=True)
ws.cell(s0 + 1, 2, "Kriterium „Systemarchitektur Multi-Site“ nach Vorlesung (Martin/Kurz 2024; Kurz 2026): Multi-Site-Systeme führen IC-Aufträge wie interne Aufträge, Mehrmandanten-Systeme über Belegaustausch zwischen Mandanten.").font = Font(italic=True)
s0 += 3
ws.cell(s0, 2, "Sensitivitätsanalyse – Nutzwert bei alternativen Gewichtungen").font = Font(bold=True, color=NAVY)
ws.cell(s0 + 1, 2, "Szenario")
for j, s_ in enumerate(D.SYSTEME):
    ws.cell(s0 + 1, 4 + NS + j, s_)
for c in ws[s0 + 1]:
    if c.value:
        c.font, c.fill = HFONT, HFILL
for k, (name, wts) in enumerate(D.SENS.items()):
    r = s0 + 2 + k
    ws.cell(r, 2, name + "  (Gewichte: " + "/".join(map(str, wts)) + ")")
    for j, col in enumerate(PC):
        terms = "+".join(f"{w}*{col}{i+2}" for i, w in enumerate(wts))
        ws.cell(r, 4 + NS + j, f"=({terms})/100").number_format = "0.00"
ws.cell(s0 + 7, 2, "Ergebnis: IFS führt funktional in allen Szenarien, proALPHA ist stabil auf Platz 2. BC holt nur bei Fokus Zeit & Pilsen auf, "
        "hat aber das beste Verhältnis von Nutzen zu Kosten. → Rangfolge Nutzen robust, Entscheidung hängt an Kosten und Zeit.").alignment = WRAP

# ------------------------------------------------------------ 7 TCO-Annahmen
ws = wb.create_sheet("7_TCO_Annahmen")
ws.append(["Parameter", "Einheit"] + D.TCO_SYS + ["Quelle / Annahme"])
for name, unit, vals, src in D.TCO_PARAM:
    ws.append([name, unit, *vals, src])
for c in ws[1]:
    c.font, c.fill, c.alignment, c.border = HFONT, HFILL, Alignment(wrap_text=True, vertical="center"), BORDER
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment, c.border = WRAP, BORDER
    for c in row[2:2 + len(D.TCO_SYS)]:
        c.fill = INFILL
        c.number_format = "0%" if row[1].value == "%" else "#,##0.##"
for i, w in enumerate([44, 16] + [16] * len(D.TCO_SYS) + [70], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "C2"
P = {name: i + 2 for i, (name, *_r) in enumerate(D.TCO_PARAM)}  # Zeilennummern


def a(name, col):
    return f"'7_TCO_Annahmen'!{col}{P[name]}"


# ------------------------------------------------------------ 8 TCO 5J
ws = wb.create_sheet("8_TCO_5J")
ws.append(["Kostenblock (T€ = Formeln aus Blatt 7)", "Erläuterung"] + D.TCO_SYS)
COLS = "".join(get_column_letter(3 + k) for k in range(len(D.TCO_SYS)))
items = [
    ("Subskription/Lizenz p. a. (Vollausbau)", "Nutzer × Preis × 12 + Hosting",
     lambda c: f"=({a('Vollnutzer',c)}*({a('Preis Vollnutzer',c)}+{a('Branchen-Add-ons (Konfigurator, APS, QM, Service mobil)',c)})+{a('Light-/Team-User',c)}*{a('Preis Light-User',c)}+{a('Shopfloor-Geräte',c)}*{a('Preis Gerät',c)})*12+{a('Hosting Private Cloud',c)}"),
    ("Einmalkosten Einführung", "Beratung + intern + Migration + Schnittstellen + Hardware + Schulung",
     lambda c: f"={a('Externe Beratertage',c)}*{a('Tagessatz extern',c)}+{a('Interne Aufwände (Key User, IT)',c)}+{a('Datenbereinigung & Migration (extern)',c)}+{a('Schnittstellen (PDM, Salesforce, EDI, Kardex, Carrier)',c)}+{a('Hardware Shopfloor/Scanner (BDE-Ersatz)',c)}+{a('Schulung & Change (inkl. BR-Forderung)',c)}"),
    ("davon externe Beratung", "PT × Tagessatz", lambda c: f"={a('Externe Beratertage',c)}*{a('Tagessatz extern',c)}"),
    ("Risikopuffer", "% auf Einmalkosten", lambda c: f"={c}3*{a('Risikopuffer auf Einmalkosten',c)}"),
    ("Subskription 5 Jahre", "Jahr 1 anteilig, Jahre 2–5 voll", lambda c: f"={c}2*({a('Subskriptionsanteil Jahr 1 (2027)',c)}+4)"),
    ("Betrieb 5 Jahre (AMS + 2. Leitung)", "AMS Jahr 1 zu 50 %", lambda c: f"={a('Betrieb/Support Partner (AMS)',c)}*4.5+{a('2. Internetleitung Singen',c)}*5"),
    ("Einmalkosten inkl. Puffer", "", lambda c: f"={c}3+{c}5"),
    ("TCO 5 Jahre (2027–2031)", "Summe", lambda c: f"={c}6+{c}7+{c}8"),
]
for lab, exp, f in items:
    ws.append([lab, exp] + [f(c) for c in COLS])
ws.append(["TCO in % vom Umsatz p. a.", "Basis 82,7 Mio. € konsolidiert"] + [f"={c}9/5/82700000" for c in COLS])
ws.append(["Abweichung zum Projektauftrag (1,2 Mio. €)", "nicht gleiche Bezugsgröße – nur Orientierung"] + [f"={c}9-1200000" for c in COLS])
ws.append(["Abweichung zur CFO-Obergrenze (1,5 Mio. € TCO)", "gleiche Bezugsgröße"] + [f"={c}9-1500000" for c in COLS])
ws.append(["Kosten in der Logik des Projektauftrags", "Einführung inkl. Puffer + 3 Jahre Lizenzen (Jahr 1 anteilig) – vergleichbar mit 1,2 Mio. €"] + [f"={c}8+{c}2*({a('Subskriptionsanteil Jahr 1 (2027)',c)}+2)" for c in COLS])
ws.append(["Vollkosten je Nutzer und Monat", "TCO / (230 Named User × 60 Monate)"] + [f"={c}9/(230*60)" for c in COLS])
ws.append(["TCO-Spanne niedrig", "TCO × Faktor niedrig (Blatt 7)"] + [f"={c}9*{a('Spanne niedrig (Faktor)',c)}" for c in COLS])
ws.append(["TCO-Spanne hoch", "TCO × Faktor hoch (Blatt 7)"] + [f"={c}9*{a('Spanne hoch (Faktor)',c)}" for c in COLS])
ws.append([])
# Jahresverlauf
ws.append(["Jahresverlauf (€)", "Jahr"] + D.TCO_SYS)
yr0 = ws.max_row
for y in range(1, 6):
    cells = []
    for c in COLS:
        sub = f"{c}2*{a('Subskriptionsanteil Jahr 1 (2027)',c)}" if y == 1 else f"{c}2"
        e = f"{c}8*{a('Anteil Einmalkosten in Jahr 1',c)}" if y == 1 else f"{c}8*(1-{a('Anteil Einmalkosten in Jahr 1',c)})" if y == 2 else "0"
        b = f"{a('Betrieb/Support Partner (AMS)',c)}*{0.5 if y == 1 else 1}+{a('2. Internetleitung Singen',c)}"
        cells.append(f"={sub}+{e}+{b}")
    ws.append([f"Jahr {y}", 2026 + y] + cells)
ws.append(["Kontrolle Summe Jahre", ""] + [f"=SUM({c}{yr0+1}:{c}{yr0+5})" for c in COLS])
for r in (1, yr0):
    for c in ws[r]:
        c.font, c.fill, c.border = HFONT, HFILL, BORDER
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment, c.border = WRAP, BORDER
        if c.column >= 3:
            c.number_format = "#,##0 €"
for r in (10,):
    for c in ws[r][2:]:
        c.number_format = "0.00%"
for c in ws[9]:
    c.font, c.fill = Font(bold=True), SUBFILL
for i, w in enumerate([44, 48] + [20] * len(D.TCO_SYS), 1):
    ws.column_dimensions[get_column_letter(i)].width = w
n0 = ws.max_row + 2
for j, t in enumerate([
        "Nicht enthalten: Weiterbetrieb KVA-WARE (Leasing Power9 bis 06/2029, Nur-Lese-Archiv), Salesforce-Verlängerung 03/2027–03/2028 (~68 T€), E-Rechnungs-Interimslösung 2027 (Annahme 15–25 T€), Lessons-Learned-/Change-Begleitung durch Dritte.",
        "Ergebnis: Keine realistische Option erreicht 1,2 Mio. € (Projektauftrag) oder < 1 Mio. € (CEO). Die CFO-Obergrenze 1,5 Mio. € TCO wird um 0,7–2,3 Mio. € überschritten (Spanne niedrig/hoch: 1,9–5,1 Mio. €) → Budgetrahmen ist GF-/Beiratsentscheidung.",
        "Einordnung: 2,2–3,8 Mio. € über 5 Jahre = 0,5–0,9 % vom Umsatz p. a. (IT-Kosten heute 1,6 % ohne Key-User- und Excel-Aufwände)."]):
    ws.cell(n0 + j, 1, t).font = Font(bold=(j == 1))

# ------------------------------------------------------------ 9 Nutzen
ws = sheet("9_Nutzen", ["Nutzenhebel", "min. T€ p. a.", "max. T€ p. a.", "Quelle", "Bemerkung"], [56, 13, 13, 40, 40], D.NUTZEN)
n = len(D.NUTZEN) + 1
ws.append(["Summe p. a. (ab Vollbetrieb 2029)", f"=SUM(B2:B{n})", f"=SUM(C2:C{n})", "", "ohne einmaligen Liquiditätseffekt Bestand 3–4 Mio. €"])
ws.append(["Kumuliert 2029–2031 (3 Jahre)", f"=B{n+1}*3", f"=C{n+1}*3", "", "Amortisation BC ca. 2031–2033, proALPHA/abas/IFS später; zusätzlich einmalig 3–4 Mio. € Liquidität aus Bestandsabbau"])
for c in ws[n + 1] + ws[n + 2]:
    c.font, c.fill, c.border = Font(bold=True), SUBFILL, BORDER
ws.append(["Nicht quantifiziert: Risiko „Do nothing“ – Ausfall KVA-WARE nach Ruhestand Brugger (1 Woche Stillstand ≈ 1,6 Mio. € Umsatz), Ausfall NAV Pilsen (40 % der Gehäuse), Strafen E-Rechnung/Exportkontrolle."])

# ------------------------------------------------------------ 10 Quellen
sheet("10_Quellen", ["ID", "Quelle (Harvard-Stil)", "Art / Prüfstatus", "Verwendet für"], [6, 100, 26, 55], D.QUELLEN)

sheet("11_Störinformationen", ["Nr", "Information", "Fundstelle", "Warum irreführend / irrelevant", "Umgang"], [4, 55, 26, 45, 50], D.STOER)
sheet("12_Lizenzmodelle", ["System", "Lizenzmodell", "Nutzerarten", "Shopfloor/Lager", "Kauf möglich?", "Konsequenz für KVA", "Prüfstatus"], [22, 40, 38, 28, 22, 45, 30], D.LIZENZ)
sheet("13_Länder_Recht", ["Gesellschaft", "Status", "Rechtliche Anforderungen an das ERP", "Weitere Anforderungen"], [24, 18, 80, 55], D.LAENDER)
sheet("14_Stakeholder", ["Person / Gremium", "Anliegen", "Antwort in unserer Empfehlung"], [28, 60, 70], D.STAKE)
sheet("15_Testskript", ["Test", "Szenario (gleich für alle Anbieter)", "Prüft"], [24, 100, 30], D.TESTS)
wb.save(OUT)
print("ok", OUT)
