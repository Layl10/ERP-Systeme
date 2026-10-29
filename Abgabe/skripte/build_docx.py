# -*- coding: utf-8 -*-
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(1.5)
sec.left_margin = sec.right_margin = Cm(1.8)
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(9.5)
st.paragraph_format.space_after = Pt(2)

h = doc.add_paragraph()
r = h.add_run("Artefakt D – KI-Nutzungsprotokoll · Fallstudie KVA ERP 2028")
r.bold, r.font.size, r.font.color.rgb = True, Pt(14), RGBColor(0x1F, 0x4E, 0x79)
doc.add_paragraph("Gruppe: [Namen / Matrikelnummern] · Stand 29.10.2026 · Vor Abgabe vom Team zu ergänzen und zu prüfen (Abschnitt 4)")

def head(t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11), RGBColor(0x1F, 0x4E, 0x79)

head("1  Eingesetzte Tools")
doc.add_paragraph("Claude (Anthropic) in Claude Code: Lesen des Fallpakets (Word, PDF, Bilder, XLSX, CSV), Strukturierung, Websuche, "
                  "Erzeugen von XLSX/PPTX per Python-Skript (openpyxl, python-pptx). LibreOffice zum Nachrechnen der Formeln und zur PDF-Ausgabe.")

head("2  Wichtigste Prompts")
t = doc.add_table(rows=1, cols=2)
t.style = "Light Grid Accent 1"
t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Prompt (sinngemäß)", "Ergebnis / Verwendung"
for a, b in [
    ("„Kannst du das mal lesen, was man alles braucht?“", "Zusammenfassung der Aufgabenstellung (Artefakte A–D, Bewertung, Termine)"),
    ("„Kannst du erstmal ein Konzept machen, wie man das machen soll?“", "Vorgehen in 5 Phasen; Struktur der Excel-Blätter"),
    ("„Kannst du das einmal komplett lösen?“", "Widerspruchsliste, Klassifizierung aller 111 Anforderungen, Longlist/Shortlist, NWA, TCO, Folien"),
    ("Anmerkungen des Dozenten (Name/Matrikelnr., gescheitertes Projekt, Störinfos, Lizenzen, Longlist/Shortlist für Rückfragen, Gesamtkosten, Management abholen, EU-Ausland/Budget)", "Ergänzt: Folie 1 (2016, Gruppe), Stakeholder-Tabelle Folie 5, Anhang A7–A10, Blätter 11–14, Sprechernotizen"),
    ("Websuchen u. a.: „E-Rechnung Pflicht Versand 2027 800.000“, „Business Central Czech localization“, „IFS Czech partner“, „proALPHA Tschechien“, „SAP S/4HANA Public Czech e-invoicing“, „NIS2UmsuCG in Kraft“",
     "Prüfung von Herstelleraussagen und Rechtsfristen gegen Primär-/Sekundärquellen"),
]:
    c = t.add_row().cells
    c[0].text, c[1].text = a, b

head("3  Wie die Ergebnisse geprüft wurden")
for s in [
    "Widersprüche nicht glätten: Jede Widerspruchszeile enthält BEIDE Aussagen wörtlich mit Fundstelle (Dokument, Interview-Nr., Datum) – Gegenprobe am Originaldokument.",
    "Vollständigkeit: Skript prüft, dass alle 111 IDs der Rohfassung klassifiziert sind (keine fehlt, keine erfunden).",
    "Rechenlogik: TCO und Nutzwert einmal in Python und einmal als Excel-Formel berechnet; LibreOffice-Neuberechnung ergab identische Werte, keine Formelfehler.",
    "Herstellerangaben: per Websuche geprüft. Bestätigt: CZ-Lokalisierung Business Central (Microsoft Learn), BC-Preise ab 11/2025 (nur USD-Sekundärquellen), IFS-Partner in CZ, SAP-KBA zu fehlendem E-Invoicing CZ. "
    "NICHT bestätigt: CZ-Präsenz von proALPHA/abas, CZ-Lokalisierung von APplus → als „offen/im RFI prüfen“ markiert statt als Fakt.",
    "Korrektur durch Recherche: Fallpaket nennt E-Rechnungs-Versandpflicht „2027/2028 offen“ – Recherche ergab 01.01.2027 für KVA (Umsatz > 800 T€) → neue Sofortmaßnahme.",
    "Preise APplus/IFS sind nicht öffentlich: alle Werte als ANNAHME (gelb) gekennzeichnet; keine von der KI „erfundenen“ Listenpreise.",
]:
    doc.add_paragraph(s, style="List Bullet")

head("4  Grenzen und Prüfschritte, die das Team selbst erledigen muss")
for s in [
    "Punktvergabe der Nutzwertanalyse ist eine begründete Einschätzung, keine Messung – im Team diskutieren und ggf. anpassen.",
    "EUR-Preise für Business Central bei Microsoft bzw. einem CSP-Partner nachschlagen; BMF-Schreiben zur E-Rechnung im Original prüfen.",
    "NIS2-Betroffenheit (Annahme: direkt betroffen als Hersteller mit ≥ 250 MA) ist keine Rechtsberatung.",
    "Folien und Excel vor der Präsentation selbst durchgehen; jede Zahl muss auf Rückfrage erklärbar sein. Rückfragen an die GF im Labor (Anhang A6) stellen.",
]:
    doc.add_paragraph(s, style="List Bullet")

doc.save(sys.argv[1])
print("ok")
